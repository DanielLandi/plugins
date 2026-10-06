import io

import pytest
from PIL import Image

from lv import ingest
from lv.util import UserError, read_text


def png(color, size=(400, 300)):
    b = io.BytesIO()
    Image.new("RGB", size, color).save(b, "PNG")
    b.seek(0)
    return b


def make_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Osmosis – ¿qué es?"
    s1.placeholders[1].text = "Water follows salt 🌊 “always”"
    s1.shapes.add_picture(png("red"), Inches(5), Inches(2))
    s1.notes_slide.notes_text_frame.text = "Ask: which way does water move?"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Data"
    rows = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    rows.cell(0, 0).text, rows.cell(0, 1).text = "Before", "After"
    rows.cell(1, 0).text, rows.cell(1, 1).text = "39 g", "43 g"
    s2.shapes.add_picture(png("red"), Inches(5), Inches(2))      # same image again: deduplicated
    s2.shapes.add_picture(png("blue"), Inches(1), Inches(4))
    prs.save(str(path))


def test_ingest_unicode_text(tmp_path):
    deck = tmp_path / "deck.pptx"
    make_pptx(deck)
    out = ingest.ingest(deck, tmp_path / "work")
    md = read_text(out / "slides.md")
    assert "## Slide 1: Osmosis – ¿qué es?" in md and "🌊 “always”" in md
    assert "Speaker notes: Ask: which way does water move?" in md
    assert "Before | After" in md and "39 g | 43 g" in md
    assert len(list((out / "media").iterdir())) == 2
    assert "Images: s001-1.png" in md and "Images: s001-1.png, s002-2.png" in md
    assert (out / "contact-01.jpg").exists()


def test_ingest_path_with_spaces_and_accents(tmp_path):
    folder = tmp_path / "Unidad 1 – Biología"
    folder.mkdir()
    deck = folder / "Clase 1 – Ósmosis.pptx"
    make_pptx(deck)
    out = ingest.ingest(deck, ingest.default_work(deck))
    assert out == folder / "lesson-videos" / "_work" / "deck" and (out / "slides.md").exists()


def test_ingest_pdf(tmp_path):
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Photosynthesis\nProducers make glucose")
    page.insert_image(pymupdf.Rect(72, 120, 272, 270), stream=png("green").getvalue())
    deck = tmp_path / "deck.pdf"
    doc.save(str(deck))
    out = ingest.ingest(deck, tmp_path / "work")
    md = read_text(out / "slides.md")
    assert "## Slide 1: Photosynthesis" in md and "Producers make glucose" in md
    assert (out / "pages" / "p001.png").exists() and len(list((out / "media").iterdir())) == 1


@pytest.mark.parametrize("ext", [".key", ".ppt", ".odp"])
def test_unsupported_formats_explain_export(tmp_path, ext):
    deck = tmp_path / f"deck{ext}"
    deck.write_bytes(b"x")
    with pytest.raises(UserError) as e:
        ingest.ingest(deck, tmp_path / "work")
    assert "PowerPoint (.pptx)" in str(e.value)


def test_contact_sheet_survives_unreadable_image(tmp_path):
    (tmp_path / "a.emf").write_bytes(b"not an image at all")
    Image.new("RGB", (50, 50), "red").save(tmp_path / "b.png")
    sheets = ingest.contact_sheets(sorted(tmp_path.iterdir()), tmp_path)
    assert len(sheets) == 1 and sheets[0].exists()
