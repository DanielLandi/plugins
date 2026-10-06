"""Read a .pptx or .pdf deck into slides.md, extracted images and labeled contact sheets for Claude to read."""
from __future__ import annotations

import hashlib
import math
from pathlib import Path

from .util import UserError, write_text

EXPORT_HELP = ("export it as PowerPoint (.pptx) first. Keynote: File → Export To → PowerPoint. "
               "Google Slides: File → Download → Microsoft PowerPoint (.pptx). Old .ppt or .odp: open it and Save As .pptx.")


def default_work(deck: Path) -> Path:
    return Path(deck).resolve().parent / "lesson-videos" / "_work"


def ingest(deck: Path, work: Path) -> Path:
    deck = Path(deck)
    ext = deck.suffix.lower()
    if ext in (".key", ".ppt", ".odp", ".gslides"):
        raise UserError(f"{deck.name}: {EXPORT_HELP}")
    if ext not in (".pptx", ".pdf"):
        raise UserError(f"{deck.name}: only .pptx and .pdf decks are supported; {EXPORT_HELP}")
    if not deck.exists():
        raise UserError(f"Can't find {deck}.")
    out = Path(work) / "deck"
    media = out / "media"
    media.mkdir(parents=True, exist_ok=True)
    slides = _pptx(deck, media) if ext == ".pptx" else _pdf(deck, out)
    write_text(out / "slides.md", _markdown(deck.name, slides))
    images = sorted(media.iterdir())
    sheets = contact_sheets(images, out)
    print(f"{len(slides)} slides, {len(images)} unique images, {len(sheets)} contact sheet(s) → {out}")
    return out


def _walk(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        try:
            group = sh.shape_type == MSO_SHAPE_TYPE.GROUP
        except Exception:  # noqa: BLE001  shapes python-pptx can't classify
            group = False
        if group:
            yield from _walk(sh.shapes)
        else:
            yield sh


def _image_blob(sh):
    try:
        im = sh.image
    except (AttributeError, ValueError, KeyError):
        return None
    return im.blob, (im.ext or "bin").lower()


def _pptx(deck: Path, media: Path) -> list[dict]:
    from pptx import Presentation
    prs = Presentation(str(deck))
    seen: dict[str, str] = {}
    slides = []
    for n, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        texts, images = [], []
        for sh in _walk(slide.shapes):
            if title_shape is not None and sh == title_shape:
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    texts.append(" | ".join(c.text.strip() for c in row.cells))
            blob = _image_blob(sh)
            if blob:
                data, ext = blob
                h = hashlib.sha1(data).hexdigest()
                if h not in seen:
                    seen[h] = f"s{n:03d}-{len(images) + 1}.{ext}"
                    (media / seen[h]).write_bytes(data)
                images.append(seen[h])
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append({"n": n, "title": title, "text": texts, "notes": notes, "images": images})
    return slides


def _pdf(deck: Path, out: Path) -> list[dict]:
    import pymupdf
    doc = pymupdf.open(str(deck))
    media, pages = out / "media", out / "pages"
    pages.mkdir(exist_ok=True)
    seen: dict[str, str] = {}
    slides = []
    for n, page in enumerate(doc, 1):
        lines = [l.strip() for l in page.get_text().splitlines() if l.strip()]
        page.get_pixmap(dpi=96).save(str(pages / f"p{n:03d}.png"))
        images = []
        for k, info in enumerate(page.get_images(full=True), 1):
            try:
                im = doc.extract_image(info[0])
            except Exception:  # noqa: BLE001
                continue
            if not im or im.get("width", 0) < 64 or im.get("height", 0) < 64:
                continue
            h = hashlib.sha1(im["image"]).hexdigest()
            if h not in seen:
                seen[h] = f"p{n:03d}-{k}.{im['ext']}"
                (media / seen[h]).write_bytes(im["image"])
            images.append(seen[h])
        slides.append({"n": n, "title": lines[0] if lines else "", "text": lines[1:], "notes": "", "images": images})
    return slides


def _markdown(name: str, slides: list[dict]) -> str:
    out = [f"# {name}", "", f"{len(slides)} slides. Images are in media/; look at contact-*.jpg to see them.", ""]
    for s in slides:
        out += [f"## Slide {s['n']}: {s['title'] or '(no title)'}", ""]
        out += s["text"]
        if s["notes"]:
            out += ["", f"Speaker notes: {s['notes']}"]
        if s["images"]:
            out += ["", "Images: " + ", ".join(s["images"])]
        out.append("")
    return "\n".join(out)


def contact_sheets(files: list[Path], out: Path, cols: int = 5, rows: int = 4) -> list[Path]:
    from PIL import Image, ImageDraw, ImageOps
    tiles = []
    for f in files:
        try:
            with Image.open(f) as im:
                im = ImageOps.exif_transpose(im).convert("RGB")
                w, h = im.size
                im.thumbnail((300, 225))
                tiles.append((f.name, f"{w}x{h}", im.copy()))
        except Exception:  # noqa: BLE001  EMF/WMF, corrupt files
            tiles.append((f.name, "no preview", None))
    per, sheets = cols * rows, []
    for s in range(math.ceil(len(tiles) / per)):
        sheet = Image.new("RGB", (cols * 320, rows * 270), "white")
        d = ImageDraw.Draw(sheet)
        for i, (name, size, im) in enumerate(tiles[s * per:(s + 1) * per]):
            x, y = (i % cols) * 320 + 10, (i // cols) * 270 + 8
            if im is not None:
                sheet.paste(im, (x + (300 - im.width) // 2, y + (225 - im.height) // 2))
            else:
                d.rectangle([x, y, x + 300, y + 225], outline="gray")
            d.text((x, y + 230), f"{name}  {size}", fill="black")
        p = out / f"contact-{s + 1:02d}.jpg"
        sheet.save(p, quality=85)
        sheets.append(p)
    return sheets
