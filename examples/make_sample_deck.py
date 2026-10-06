# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx>=1.0", "pillow>=10.0"]
# ///
"""Build examples/water-cycle.pptx: a 6-slide, photo-free sample deck for trying lesson-videos."""
import io
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

OUT = Path(__file__).resolve().parent / "water-cycle.pptx"


def drawing(kind: str) -> io.BytesIO:
    im = Image.new("RGB", (1200, 800), "#DFF3FB")
    d = ImageDraw.Draw(im)
    d.rectangle([0, 620, 1200, 800], fill="#2E86C1")
    if kind in ("sun", "cycle"):
        d.ellipse([920, 60, 1120, 260], fill="#F7C531")
    if kind in ("cloud", "cycle"):
        for x, y, r in [(380, 220, 90), (480, 180, 110), (600, 220, 90)]:
            d.ellipse([x - r, y - r, x + r, y + r], fill="white")
    if kind == "cycle":
        d.line([(900, 560), (700, 300)], fill="#1B4F72", width=12)
        d.line([(480, 330), (480, 600)], fill="#1B4F72", width=12)
    b = io.BytesIO()
    im.save(b, "PNG")
    b.seek(0)
    return b


def main():
    prs = Presentation()
    slides = [
        ("The Water Cycle", "Unit 3 · How water moves around Earth", None, "Essential question: How does water move between the ocean, the air and the land?"),
        ("Evaporation", "The sun heats water in oceans and lakes.\nLiquid water turns into water vapor, an invisible gas.", "sun", "Key word: evaporation = liquid to gas."),
        ("Condensation", "Water vapor cools high in the atmosphere.\nIt turns back into tiny liquid droplets.\nMany droplets together form clouds.", "cloud", "Common mistake: clouds are liquid droplets, not water vapor."),
        ("Precipitation", "Droplets join and get heavy.\nThey fall as rain, snow, sleet or hail.", None, ""),
        ("Collection", "Water collects in oceans, lakes, rivers and underground.\nThen the cycle starts again.", "cycle", ""),
        ("Review", "1. Name the three main steps of the water cycle.\n2. Which step turns liquid into gas?\n3. Are clouds made of gas or liquid?", None, "Answers: evaporation, condensation, precipitation; evaporation; liquid droplets."),
    ]
    for title, body, pic, notes in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
        if pic:
            s.shapes.add_picture(drawing(pic), Inches(5.2), Inches(2.2), width=Inches(4.3))
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    prs.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
